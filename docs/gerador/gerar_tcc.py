"""Gera docs/TCC-EBAC-QE.docx a partir do template do professor.

Uso (na raiz do repositório):
    python docs/gerador/gerar_tcc.py "<caminho do template .docx>"

O conteúdo textual fica em conteudo.py; resultados de performance são lidos de Performance/reports/*.json.
"""
import copy
import json
import os
import sys

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(__file__))
import conteudo as C  # noqa: E402

RAIZ = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
SAIDA = os.path.join(RAIZ, 'docs', 'TCC-EBAC-QE.docx')
BULLET_NUM_ID = 19
SUMARIO = [('1.', 'RESUMO'), ('2.', 'SUMÁRIO'), ('3.', 'INTRODUÇÃO'), ('4.', 'O PROJETO'),
           ('4.1', 'Estratégia de teste'), ('4.2', 'Critérios de aceitação'), ('4.3', 'Casos de testes'),
           ('4.4', 'Repositório no Github'), ('4.5', 'Testes automatizados'), ('4.6', 'Integração contínua'),
           ('4.7', 'Testes de performance'), ('5.', 'CONCLUSÃO'), ('6.', 'REFERÊNCIAS BIBLIOGRÁFICAS')]
LARGURA_UTIL = Cm(15)
_figuras = [0]


# ---------------------------------------------------------------- utilitários
def definir_texto(paragrafo, texto):
    runs = paragrafo.runs
    runs[0].text = texto
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def quebra_pagina(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def titulo(doc, texto, nivel=1, nova_pagina=False):
    if nova_pagina:
        quebra_pagina(doc)
    return doc.add_paragraph(texto, style=f'Heading {nivel}')


def texto(doc, conteudo, negrito_inicio=None, alinhamento=WD_ALIGN_PARAGRAPH.JUSTIFY, tamanho=None):
    p = doc.add_paragraph(style='Normal')
    p.alignment = alinhamento
    p.paragraph_format.first_line_indent = Cm(1.25) if alinhamento == WD_ALIGN_PARAGRAPH.JUSTIFY and not negrito_inicio else None
    p.paragraph_format.line_spacing = 1.5
    if negrito_inicio:
        r = p.add_run(negrito_inicio)
        r.bold = True
        if tamanho:
            r.font.size = tamanho
    r = p.add_run(conteudo)
    if tamanho:
        r.font.size = tamanho
    return p


def marcador(doc, conteudo, negrito_inicio=None, nivel=0):
    p = doc.add_paragraph(style='List Paragraph')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement('w:numPr')
    ilvl = OxmlElement('w:ilvl')
    ilvl.set(qn('w:val'), str(nivel))
    numId = OxmlElement('w:numId')
    numId.set(qn('w:val'), str(BULLET_NUM_ID))
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.insert(0, numPr) if pPr.find(qn('w:pStyle')) is None else pPr.find(qn('w:pStyle')).addnext(numPr)
    if negrito_inicio:
        p.add_run(negrito_inicio).bold = True
    p.add_run(conteudo)
    return p


def sombrear(elemento_pr, cor):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), cor)
    elemento_pr.append(shd)


def bloco_codigo(doc, linhas):
    for linha in linhas:
        p = doc.add_paragraph(style='Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        fmt = p.paragraph_format
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0
        fmt.left_indent = Cm(0)
        sombrear(p._p.get_or_add_pPr(), 'F2F2F2')
        r = p.add_run(linha if linha else ' ')
        r.font.name = 'Consolas'
        r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Consolas')
        r.font.size = Pt(7.5)
        if linha.strip().startswith(('Funcionalidade', 'Cenário', 'Esquema do Cenário', 'Contexto', 'Exemplos')):
            r.bold = True
            r.font.color.rgb = RGBColor(0x6B, 0x2C, 0x91)
        elif linha.strip().startswith('#'):
            r.font.color.rgb = RGBColor(0x70, 0x70, 0x70)
    doc.add_paragraph(style='Normal')


def bordas_tabela(tabela):
    tblPr = tabela._tbl.tblPr
    bordas = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement(f'w:{lado}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '808080')
        bordas.append(b)
    tblPr.append(bordas)


def tabela(doc, cabecalho, linhas, larguras_cm, tamanho=8, legenda=None):
    if legenda:
        p = doc.add_paragraph(style='Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(legenda)
        r.bold = True
        r.font.size = Pt(10)
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    bordas_tabela(t)
    larguras = [Cm(x) for x in larguras_cm]

    def preencher(celula, valor, largura, negrito=False, fundo=None):
        celula.width = largura
        if fundo:
            sombrear(celula._tc.get_or_add_tcPr(), fundo)
        p = celula.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = None
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        partes = str(valor).split('\n')
        for i, parte in enumerate(partes):
            if i:
                p = celula.add_paragraph()
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
            r = p.add_run(parte)
            r.font.size = Pt(tamanho)
            r.bold = negrito

    for i, largura in enumerate(larguras):
        t.columns[i].width = largura
    for i, cab in enumerate(cabecalho):
        preencher(t.rows[0].cells[i], cab, larguras[i], negrito=True, fundo='E4D5EE')
    for linha in linhas:
        cells = t.add_row().cells
        for i, valor in enumerate(linha):
            preencher(cells[i], valor, larguras[i])
    # cabeçalho repete em quebras de página
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement('w:tblHeader')
    th.set(qn('w:val'), 'true')
    trPr.append(th)
    doc.add_paragraph(style='Normal')
    return t


def imagem(doc, caminho, largura=LARGURA_UTIL, legenda=None):
    if not os.path.exists(caminho):
        texto(doc, f'[Imagem não encontrada: {os.path.relpath(caminho, RAIZ)}]')
        return
    p = doc.add_paragraph(style='Normal')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(caminho, width=largura)
    if legenda:
        lp = doc.add_paragraph(style='Normal')
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _figuras[0] += 1
        r = lp.add_run(f'Figura {_figuras[0]} – {legenda}')
        r.italic = True
        r.font.size = Pt(9)


def ler_feature(nome):
    with open(os.path.join(RAIZ, 'docs', 'features', nome), encoding='utf-8') as f:
        return [l.rstrip('\n') for l in f if not l.startswith('# language')]


def metricas_k6(nome):
    caminho = os.path.join(RAIZ, 'Performance', 'reports', f'{nome}.json')
    if not os.path.exists(caminho):
        return None
    with open(caminho, encoding='utf-8') as f:
        m = json.load(f)['metrics']
    v = lambda k: m.get(k, {}).get('values', {})  # noqa: E731
    return {
        'reqs': int(v('http_reqs').get('count', 0)),
        'rps': v('http_reqs').get('rate', 0),
        'iter': int(v('iterations').get('count', 0)),
        'avg': v('http_req_duration').get('avg', 0) / 1000,
        'med': v('http_req_duration').get('med', 0) / 1000,
        'p90': v('http_req_duration').get('p(90)', 0) / 1000,
        'p95': v('http_req_duration').get('p(95)', 0) / 1000,
        'max': v('http_req_duration').get('max', 0) / 1000,
        'falhas': v('http_req_failed').get('rate', 0) * 100,
        'checks': v('checks').get('rate', 0) * 100,
        'vus': int(v('vus_max').get('max', v('vus_max').get('value', 0))),
    }


# ---------------------------------------------------------------- documento
def main(template):
    doc = Document(template)
    P = doc.paragraphs

    # capa
    definir_texto(P[6], C.AUTOR)
    definir_texto(P[13], 'Análise de Qualidade – EBAC Shop')
    definir_texto(P[21], C.CIDADE)
    definir_texto(P[22], C.ANO)

    # resumo
    definir_texto(P[24], C.RESUMO)

    # sumário: mantém as entradas do template, atualizando os números de página (paginas.json)
    paginas = {}
    arq_paginas = os.path.join(os.path.dirname(__file__), 'paginas.json')
    if os.path.exists(arq_paginas):
        with open(arq_paginas, encoding='utf-8') as f:
            paginas = json.load(f)
    for p, (num, nome) in zip(P[28:41], SUMARIO):
        for filho in list(p._p):
            if filho.tag != qn('w:pPr'):
                p._p.remove(filho)
        p.add_run(f'{num}	{nome}	{paginas.get(num, "")}')

    # remove todo o conteúdo-modelo a partir da introdução (mantém a quebra de página após o sumário)
    for p in P[42:]:
        p._p.getparent().remove(p._p)

    # ---------------- INTRODUÇÃO
    titulo(doc, 'INTRODUÇÃO', 1, nova_pagina=True)
    for par in C.INTRODUCAO:
        texto(doc, par)

    # ---------------- O PROJETO
    titulo(doc, 'O PROJETO', 1, nova_pagina=True)
    for par in C.PROJETO:
        texto(doc, par)
    for item in C.PROJETO_AMBIENTES:
        marcador(doc, item[1], item[0])

    # 4.1 Estratégia
    titulo(doc, 'Estratégia de teste', 2)
    for par in C.ESTRATEGIA_TEXTO:
        texto(doc, par)
    tabela(doc, ['Diretriz', 'Definição'], C.ESTRATEGIA_TABELA, [3.6, 11.4], tamanho=8.5,
           legenda='Tabela 1 – Resumo da estratégia de testes')

    # mapa mental em página paisagem
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = sec.page_height, sec.page_width
    p = doc.add_paragraph(style='Normal')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    imagem(doc, os.path.join(RAIZ, 'docs', 'estrategia-mapa-mental.png'),
           largura=sec.page_width - sec.left_margin - sec.right_margin,
           legenda='Mapa mental da estratégia de testes. Fonte: elaborado pelo autor.')
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = sec.page_height, sec.page_width

    # 4.2 Critérios de aceitação
    titulo(doc, 'Critérios de aceitação', 2)
    texto(doc, C.CRITERIOS_INTRO)
    for us in C.HISTORIAS:
        p = doc.add_paragraph(style='Normal')
        r = p.add_run(us['titulo'])
        r.bold = True
        r.font.color.rgb = RGBColor(0x6B, 0x2C, 0x91)
        if us.get('nova'):
            for rotulo, valor in (('Como ', us['como']), ('Quero ', us['quero']), ('Para ', us['para'])):
                texto(doc, valor, rotulo, alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
            texto(doc, '', 'Regras de negócio:', alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
            for regra in us['regras']:
                marcador(doc, regra)
        texto(doc, '', 'Critérios de aceitação (Gherkin):', alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
        linhas = [l for l in ler_feature(us['feature']) if not (us.get('nova') and l.strip().startswith('#'))]
        bloco_codigo(doc, linhas)

    # 4.3 Casos de teste
    titulo(doc, 'Casos de testes', 2)
    for par in C.CASOS_INTRO:
        texto(doc, par)
    for i, grupo in enumerate(C.CASOS, start=2):
        tabela(doc, ['ID', 'Caso de teste / passos', 'Técnica', 'Resultado esperado', 'Fluxo', 'Autom.'],
               grupo['casos'], [2.0, 4.4, 2.0, 3.8, 1.7, 1.1], tamanho=7.5,
               legenda=f'Tabela {i} – Casos de teste {grupo["us"]}')
    texto(doc, C.DEFEITOS_INTRO)
    tabela(doc, ['ID', 'Defeito', 'Evidência', 'Severidade'], C.DEFEITOS, [1.6, 5.2, 6.4, 1.8], tamanho=8,
           legenda=f'Tabela {len(C.CASOS) + 2} – Defeitos encontrados durante a execução')

    # 4.4 Repositório
    titulo(doc, 'Repositório no Github', 2)
    texto(doc, C.REPOSITORIO_TEXTO)
    texto(doc, C.REPOSITORIO_LINK, 'Link do repositório: ', alinhamento=WD_ALIGN_PARAGRAPH.LEFT)
    bloco_codigo(doc, C.REPOSITORIO_ARVORE)

    # 4.5 Testes automatizados
    titulo(doc, 'Testes automatizados', 2)
    texto(doc, C.AUTOMACAO_INTRO)
    p = doc.add_paragraph('Automação de UI', style='Heading 3')
    for par in C.UI_TEXTO:
        texto(doc, par)
    tabela(doc, ['Critério', 'Cypress (JavaScript)', 'Playwright (TypeScript)', 'Selenium WebDriver (Java)'],
           C.COMPARATIVO, [3.0, 4.0, 4.0, 4.0], tamanho=8,
           legenda=f'Tabela {len(C.CASOS) + 3} – Comparativo de ferramentas de automação web')
    for par in C.UI_JUSTIFICATIVA:
        texto(doc, par)
    imagem(doc, os.path.join(RAIZ, 'docs', 'evidencias', 'ui-report.png'),
           legenda='Relatório Mochawesome da automação de UI (9 de 9 testes aprovados).')

    doc.add_paragraph('Automação de API', style='Heading 3')
    for par in C.API_TEXTO:
        texto(doc, par)
    imagem(doc, os.path.join(RAIZ, 'docs', 'evidencias', 'api-report.png'),
           legenda='Relatório Mochawesome da automação de API (8 de 8 testes aprovados).')

    doc.add_paragraph('Automação Mobile', style='Heading 3')
    for par in C.MOBILE_TEXTO:
        texto(doc, par)
    for nome, legenda in C.MOBILE_IMAGENS:
        caminho = os.path.join(RAIZ, 'docs', 'evidencias', nome)
        if os.path.exists(caminho):
            imagem(doc, caminho, largura=Cm(6), legenda=legenda)

    # 4.6 CI
    titulo(doc, 'Integração contínua', 2)
    for par in C.CI_TEXTO:
        texto(doc, par)
    for item in C.CI_ITENS:
        marcador(doc, item[1], item[0])
    imagem(doc, os.path.join(RAIZ, 'docs', 'evidencias', 'github-actions.png'),
           legenda='Execução do pipeline no GitHub Actions.')

    # 4.7 Performance
    titulo(doc, 'Testes de performance', 2)
    for par in C.PERF_TEXTO:
        texto(doc, par)
    linhas = []
    for nome, rotulo in (('login', 'CT-PERF-01 Login'), ('catalogo', 'CT-PERF-02 Catálogo')):
        m = metricas_k6(nome)
        if m:
            linhas.append([rotulo, m['vus'], m['reqs'], f"{m['rps']:.2f}", f"{m['avg']:.2f}", f"{m['p95']:.2f}",
                           f"{m['max']:.2f}", f"{m['falhas']:.2f}%", f"{m['checks']:.2f}%"])
    tabela(doc, ['Cenário', 'VUs', 'Requisições', 'Req/s', 'Média (s)', 'p95 (s)', 'Máx (s)', 'Falhas', 'Checks'],
           linhas, [2.9, 1.0, 1.7, 1.3, 1.6, 1.4, 1.4, 1.4, 1.4], tamanho=8,
           legenda=f'Tabela {len(C.CASOS) + 4} – Resultados dos testes de carga (20 VUs, 2 min, ramp-up 20 s)')
    for par in C.PERF_ANALISE:
        texto(doc, par)
    imagem(doc, os.path.join(RAIZ, 'docs', 'evidencias', 'k6-login.png'),
           legenda='Relatório k6 do cenário de login.')
    imagem(doc, os.path.join(RAIZ, 'docs', 'evidencias', 'k6-catalogo.png'),
           legenda='Relatório k6 do cenário de catálogo.')

    # ---------------- CONCLUSÃO
    titulo(doc, 'CONCLUSÃO', 1, nova_pagina=True)
    for par in C.CONCLUSAO:
        texto(doc, par)

    # ---------------- REFERÊNCIAS
    titulo(doc, 'REFERÊNCIAS BIBLIOGRÁFICAS', 1, nova_pagina=True)
    for ref in C.REFERENCIAS:
        p = doc.add_paragraph(style='Normal')
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.0
        for parte, negrito in ref:
            p.add_run(parte).bold = negrito

    doc.core_properties.author = C.AUTOR
    doc.core_properties.title = 'TCC - Engenheiro de Qualidade de Software'
    doc.save(SAIDA)
    print('Gerado:', SAIDA)


if __name__ == '__main__':
    main(sys.argv[1])
